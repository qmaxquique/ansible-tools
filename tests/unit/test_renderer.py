from copy import deepcopy
from datetime import date, datetime, timezone
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from plugins.module_utils import xlsx_report as renderer


@pytest.fixture
def definition():
    return {
        "styles": {
            "header": {
                "font": {"bold": True, "color": "FFFFFF", "size": 12},
                "fill": {"color": "#1F4E78"},
                "alignment": {"horizontal": "center", "wrap_text": True},
                "border": {"bottom": {"style": "thin", "color": "FFFFFF"}},
            },
            "failed": {"background": "FFC7CE", "font": {"bold": True, "color": "9C0006"}},
            "percent": {"number_format": "0.0%"},
        },
        "properties": {"title": "Infrastructure report", "last_modified_by": "Ansible"},
        "sheets": [
            {
                "name": "Summary",
                "header_style": "header",
                "columns": [
                    {"name": "Server", "width": 25},
                    {"name": "Status", "width": 18},
                    {"name": "Updates", "number_format": "0"},
                ],
                "rows": [
                    ["server01", "SUCCESS", 12],
                    {"Server": "server02", "Status": "FAILED", "Updates": 3},
                    {"values": ["server03", "FAILED", 0], "styles": {"Status": "failed"}},
                ],
                "freeze_panes": "A2",
                "auto_filter": True,
                "tables": [{"name": "Inventory"}],
                "conditional_formatting": [
                    {"column": "Status", "equals": "FAILED", "style": "failed"},
                    {"range": "A2:C4", "formula": '=$B2="FAILED"', "background": "FFC7CE"},
                ],
            },
            {
                "name": "Details",
                "columns": [{"name": "Metric"}, {"name": "Value"}],
                "rows": [
                    ["Total", {"formula": "=SUM(Summary!C2:C4)"}],
                    [
                        "Date",
                        {"value": "2026-10-01", "type": "date", "number_format": "yyyy-mm-dd"},
                    ],
                    ["Time", {"value": "2026-10-01T14:30:00", "type": "datetime"}],
                    ["Compliant", True],
                    ["Ratio", {"value": 0.95, "style": "percent"}],
                    ["Input", '=HYPERLINK("https://example.com")'],
                    ["Blank", None],
                ],
            },
        ],
    }


def render(definition):
    spec = renderer.normalize_report(**definition)
    return renderer.render_workbook(spec)


def test_roundtrip_preserves_report_features(definition):
    workbook = load_workbook(BytesIO(render(definition)))
    try:
        assert workbook.sheetnames == ["Summary", "Details"]
        summary, details = workbook.worksheets
        assert summary["A1"].value == "Server"
        assert summary["A1"].font.bold
        assert summary["A1"].fill.fgColor.rgb == "FF1F4E78"
        assert summary["A1"].alignment.horizontal == "center"
        assert summary["A1"].border.bottom.style == "thin"
        assert summary["A3"].value == "server02"
        assert summary["B4"].fill.fgColor.rgb == "FFFFC7CE"
        assert summary["C2"].value == 12 and summary["C2"].number_format == "0"
        assert summary.column_dimensions["A"].width == 25
        assert summary.freeze_panes == "A2" and summary.auto_filter.ref is None
        assert summary.tables["Inventory"].ref == "A1:C4"
        assert summary.tables["Inventory"].autoFilter.ref == "A1:C4"
        assert summary.tables["Inventory"].tableStyleInfo.showRowStripes
        rules = list(summary.conditional_formatting)
        assert str(rules[0].sqref) == "B2:B4"
        assert summary.conditional_formatting[rules[0]][0].formula == ['"FAILED"']
        assert summary.conditional_formatting[rules[0]][0].dxf.font.bold
        assert summary.conditional_formatting[rules[1]][0].formula == ['$B2="FAILED"']
        assert details["B2"].value == "=SUM(Summary!C2:C4)" and details["B2"].data_type == "f"
        assert details["B3"].value == date(2026, 10, 1)
        assert details["B4"].value == datetime(2026, 10, 1, 14, 30)
        assert details["B5"].value is True
        assert details["B6"].number_format == "0.0%"
        assert details["B7"].data_type == "s"
        assert details["B8"].value is None
        assert workbook.properties.title == "Infrastructure report"
        assert workbook.properties.lastModifiedBy == "Ansible"
    finally:
        workbook.close()


def test_render_is_reproducible_and_does_not_mutate_inputs(definition):
    original = deepcopy(definition)
    first = render(definition)
    second = render(definition)
    assert first == second
    assert renderer.workbook_checksum(first) == renderer.workbook_checksum(second)
    assert len(renderer.workbook_checksum(first)) == 64
    assert definition == original
    with ZipFile(BytesIO(first)) as archive:
        core = archive.read("docProps/core.xml")
        assert b"xmlns:dcterms=" in core
        ElementTree.fromstring(core)


@pytest.mark.parametrize("auto_filter", [True, "A1:C4", "$A$1:$C$4"])
def test_table_filter_is_not_duplicated_on_worksheet(definition, auto_filter):
    definition["sheets"][0]["auto_filter"] = auto_filter
    with ZipFile(BytesIO(render(definition))) as archive:
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        worksheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        table = ElementTree.fromstring(archive.read("xl/tables/table1.xml"))
        # Excel repairs an overlapping worksheet filter by removing the structured table.
        assert worksheet.find("s:autoFilter", namespace) is None
        assert worksheet.find("s:tableParts", namespace) is not None
        assert table.find("s:autoFilter", namespace).get("ref") == "A1:C4"
        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        assert not any(
            name.get("name") == "_xlnm._FilterDatabase"
            for name in workbook.findall("s:definedNames/s:definedName", namespace)
        )


def test_auto_filter_without_table_is_preserved(definition):
    definition["sheets"][0]["tables"] = []
    workbook = load_workbook(BytesIO(render(definition)))
    try:
        assert workbook["Summary"].auto_filter.ref == "A1:C4"
    finally:
        workbook.close()


@pytest.mark.parametrize("auto_filter", ["A1:C4", "A1:B2", "B1:C4"])
def test_explicit_filter_partially_overlapping_table_is_rejected(definition, auto_filter):
    definition["sheets"][0]["tables"][0]["ref"] = "A1:B4"
    definition["sheets"][0]["auto_filter"] = auto_filter
    with pytest.raises(renderer.ReportError, match=r"sheets\[0\].auto_filter.*overlap"):
        renderer.normalize_report(**definition)


def test_non_overlapping_worksheet_and_table_filters_can_coexist(definition):
    definition["sheets"][0]["tables"][0]["ref"] = "A1:B4"
    definition["sheets"][0]["auto_filter"] = "C1:C4"
    workbook = load_workbook(BytesIO(render(definition)))
    try:
        sheet = workbook["Summary"]
        assert sheet.auto_filter.ref == "C1:C4"
        assert sheet.tables["Inventory"].autoFilter.ref == "A1:B4"
    finally:
        workbook.close()


def test_auto_filter_true_uses_filters_of_multiple_tables(definition):
    definition["sheets"][0]["tables"] = [
        {"name": "Hosts", "ref": "A1:A4"},
        {"name": "Results", "ref": "B1:C4"},
    ]
    workbook = load_workbook(BytesIO(render(definition)))
    try:
        sheet = workbook["Summary"]
        assert sheet.auto_filter.ref is None
        assert [table.autoFilter.ref for table in sheet.tables.values()] == ["A1:A4", "B1:C4"]
    finally:
        workbook.close()


def repack(data, transform):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as destination:
        for name in reversed(source.namelist()):
            destination.writestr(name, transform(name, source.read(name)))
    return output.getvalue()


def test_checksum_ignores_save_times_and_zip_metadata(definition):
    data = render(definition)
    repacked = repack(
        data,
        lambda name, value: (
            value.replace(b"2000-01-01", b"2026-10-01") if name == "docProps/core.xml" else value
        ),
    )
    assert data != repacked
    assert renderer.workbook_checksum(data) == renderer.workbook_checksum(repacked)


@pytest.mark.parametrize(
    "component", ["cell", "style", "properties", "sheet_order", "table", "rule"]
)
def test_checksum_detects_material_changes(definition, component):
    previous = render(definition)
    if component == "cell":
        definition["sheets"][0]["rows"][0][2] = 13
    elif component == "style":
        definition["styles"]["header"]["font"]["bold"] = False
    elif component == "properties":
        definition["properties"]["title"] = "Another title"
    elif component == "sheet_order":
        definition["sheets"].reverse()
    elif component == "table":
        definition["sheets"][0]["tables"][0]["style"] = "TableStyleMedium9"
    elif component == "rule":
        definition["sheets"][0]["conditional_formatting"][0]["equals"] = "WARNING"
    assert renderer.workbook_checksum(previous) != renderer.workbook_checksum(render(definition))


def test_style_precedence_and_number_format():
    styles = {
        name: {"background": color}
        for name, color in {
            "column": "111111",
            "row": "222222",
            "override": "333333",
            "cell": "444444",
        }.items()
    }
    definition = {
        "styles": styles,
        "sheets": [
            {
                "name": "Styles",
                "columns": [{"name": "Value", "style": "column", "number_format": "0.00"}],
                "rows": [
                    [1],
                    {"values": [2], "style": "row"},
                    {"values": [3], "style": "row", "styles": {"Value": "override"}},
                    {
                        "values": [{"value": 4, "style": "cell", "number_format": "0.0"}],
                        "style": "row",
                        "styles": {"Value": "override"},
                    },
                ],
            }
        ],
    }
    workbook = load_workbook(BytesIO(render(definition)))
    try:
        assert [workbook.active.cell(row, 1).fill.fgColor.rgb for row in range(2, 6)] == [
            "FF111111",
            "FF222222",
            "FF333333",
            "FF444444",
        ]
        assert workbook.active["A2"].number_format == "0.00"
        assert workbook.active["A5"].number_format == "0.0"
    finally:
        workbook.close()


def test_empty_sheet_and_keyed_missing_cells():
    spec = renderer.normalize_report(
        [
            {
                "name": "Empty",
                "columns": [{"name": "Status"}],
                "tables": [{"name": "EmptyTable"}],
                "conditional_formatting": [
                    {"column": "Status", "equals": "OK", "background": "00FF00"}
                ],
            },
            {
                "name": "Missing",
                "columns": [{"name": "Host"}, {"name": "Count"}],
                "rows": [{"Host": "one"}],
            },
        ]
    )
    workbook = load_workbook(BytesIO(renderer.render_workbook(spec)))
    try:
        assert workbook["Empty"].tables["EmptyTable"].ref == "A1:A1"
        assert len(workbook["Empty"].conditional_formatting) == 0
        assert workbook["Missing"]["B2"].value is None
    finally:
        workbook.close()


def test_date_format_survives_a_color_only_style():
    data = render(
        {
            "styles": {"highlight": {"background": "FFEB9C"}},
            "sheets": [
                {
                    "name": "Dates",
                    "columns": [{"name": "When", "style": "highlight"}],
                    "rows": [[{"value": "2026-10-01", "type": "date"}]],
                }
            ],
        }
    )
    workbook = load_workbook(BytesIO(data))
    try:
        assert workbook.active["A2"].number_format == "yyyy-mm-dd"
        assert workbook.active["A2"].value == date(2026, 10, 1)
    finally:
        workbook.close()


def test_conditional_equality_escapes_text_and_overrides_font_color():
    data = render(
        {
            "styles": {"highlight": {"font": {"bold": True, "color": "FF0000"}}},
            "sheets": [
                {
                    "name": "Status",
                    "columns": [{"name": "Status"}],
                    "rows": [['Say "OK"']],
                    "conditional_formatting": [
                        {
                            "column": "Status",
                            "equals": 'Say "OK"',
                            "style": "highlight",
                            "font_color": "00FF00",
                        }
                    ],
                }
            ],
        }
    )
    workbook = load_workbook(BytesIO(data))
    try:
        conditional = next(iter(workbook.active.conditional_formatting))
        rule = workbook.active.conditional_formatting[conditional][0]
        assert rule.formula == ['"Say ""OK"""']
        assert rule.dxf.font.bold and rule.dxf.font.color.rgb == "FF00FF00"
    finally:
        workbook.close()


def test_style_names_must_be_unique_without_regard_to_case():
    with pytest.raises(renderer.ReportError, match="duplicate names"):
        renderer.normalize_report(
            [{"name": "Sheet", "columns": [{"name": "Host"}]}],
            {"header": {"background": "FFFFFF"}, "Header": {"background": "FFFFFF"}},
        )


@pytest.mark.parametrize(
    "value",
    [
        float("nan"),
        float("inf"),
        10**400,
        [1, 2],
        b"bytes",
        "bad\x00text",
        "x" * 32768,
        datetime(2026, 10, 1, tzinfo=timezone.utc),
        {"value": "bad", "type": "date"},
        {"value": "2026-10-01T00:00:00+00:00", "type": "datetime"},
        {"value": True, "type": "number"},
        {"value": "x", "type": []},
        {"formula": "SUM(A1)"},
        {"formula": "=1", "value": 1},
        {"value": 1, "style": "missing"},
    ],
)
def test_invalid_cells_fail_with_context(value):
    with pytest.raises(renderer.ReportError, match=r"sheets\[0\].rows\[0\]"):
        renderer.normalize_report(
            [{"name": "Sheet", "columns": [{"name": "Value"}], "rows": [[value]]}]
        )


@pytest.mark.parametrize(
    "options",
    [
        {"name": "bad/name"},
        {"name": "'quoted'"},
        {"name": "x" * 32},
        {"columns": []},
        {"columns": [{"name": "Host"}, {"name": "host"}]},
        {"columns": [{"name": "Host", "width": False}]},
        {"rows": [[1, 2]]},
        {"rows": [{"unknown": 1}]},
        {"freeze_panes": "XFE2"},
        {"freeze_panes": "A0"},
        {"auto_filter": "A1:B2"},
        {"auto_filter": "A2:A1"},
        {"tables": [{"name": "A1"}]},
        {"tables": [{"name": "R1C1"}]},
        {"tables": [{"name": "Space Name"}]},
        {"tables": [{"name": "TableOne", "ref": "A2:A2"}]},
        {"tables": [{"name": "TableOne", "style": "TableStyleMedium99"}]},
        {"tables": [{"name": "One"}, {"name": "Two"}]},
        {"conditional_formatting": [{"column": "unknown", "equals": 1, "background": "FF0000"}]},
        {"conditional_formatting": [{"column": "Host", "equals": 1}]},
        {"conditional_formatting": [{"range": "A1:A9", "equals": 1, "background": "FF0000"}]},
        {"conditional_formatting": [{"column": "Host", "formula": "=", "background": "FF0000"}]},
        {"unsupported": True},
    ],
)
def test_invalid_sheet_definitions_are_rejected(options):
    sheet = {"name": "Sheet", "columns": [{"name": "Host"}], "rows": [["one"]]}
    sheet.update(options)
    with pytest.raises(renderer.ReportError, match=r"sheets\[0\]"):
        renderer.normalize_report([sheet])


@pytest.mark.parametrize(
    "style",
    [
        {},
        {"font": {"bold": "yes"}},
        {"font": {"size": -1}},
        {"fill": {"color": "red"}},
        {"alignment": {"horizontal": []}},
        {"border": {"bottom": {"style": "invalid"}}},
        {"background": "FFFFFF", "fill": {"color": "FFFFFF"}},
        {"font": {"underline": ["single"]}},
        {"alignment": {"indent": 1.5}},
    ],
)
def test_invalid_styles_are_rejected(style):
    with pytest.raises(renderer.ReportError, match="styles.header"):
        renderer.normalize_report(
            [{"name": "Sheet", "columns": [{"name": "Host"}]}], {"header": style}
        )


def test_duplicate_sheet_and_table_names_are_rejected():
    first = {"name": "One", "columns": [{"name": "Host"}], "tables": [{"name": "Inventory"}]}
    with pytest.raises(renderer.ReportError, match="duplicates another worksheet"):
        renderer.normalize_report([first, dict(first, name="one")])
    with pytest.raises(renderer.ReportError, match="duplicates another table"):
        renderer.normalize_report([first, dict(first, name="Two")])


def test_missing_or_unsupported_dependency_is_actionable(monkeypatch):
    monkeypatch.setattr(renderer, "HAS_OPENPYXL", False)
    with pytest.raises(renderer.ReportError, match="must be installed"):
        renderer.require_openpyxl()
    monkeypatch.setattr(renderer, "HAS_OPENPYXL", True)
    monkeypatch.setattr(renderer.openpyxl, "__version__", "3.0.10")
    with pytest.raises(renderer.ReportError, match="found 3.0.10"):
        renderer.require_openpyxl()


@pytest.mark.parametrize("data", [b"", b"not a zip file", b"PK\x03\x04broken"])
def test_corrupt_existing_files_have_no_checksum(data):
    assert renderer.workbook_checksum(data) is None
