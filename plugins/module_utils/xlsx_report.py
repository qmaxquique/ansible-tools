"""Validate and render declarative XLSX workbooks without importing Ansible."""

import hashlib
import math
import re
import zlib
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile, ZipInfo

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.formatting.rule import Rule
    from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
    from openpyxl.styles.differential import DifferentialStyle
    from openpyxl.styles.numbers import NumberFormat
    from openpyxl.utils import get_column_letter, range_boundaries
    from openpyxl.worksheet.table import Table, TableStyleInfo

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


MAX_ROWS = 1048576
MAX_COLUMNS = 16384
INVALID_XML = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
CELL_REFERENCE = re.compile(r"[A-Z]{1,3}[1-9][0-9]*\Z")
RANGE_REFERENCE = re.compile(r"[A-Z]{1,3}[1-9][0-9]*(?::[A-Z]{1,3}[1-9][0-9]*)?\Z")
EPOCH = datetime(2000, 1, 1)


class ReportError(ValueError):
    """An invalid report definition or unsupported runtime dependency."""


@dataclass(frozen=True)
class CellSpec:
    value: object
    style: object = None
    number_format: object = None
    formula: bool = False


@dataclass(frozen=True)
class ColumnSpec:
    name: str
    width: object = None
    style: object = None
    header_style: object = None
    number_format: object = None


@dataclass(frozen=True)
class SheetSpec:
    name: str
    columns: tuple
    rows: tuple
    freeze_panes: object
    auto_filter: object
    tables: tuple
    conditional_formatting: tuple


@dataclass(frozen=True)
class WorkbookSpec:
    sheets: tuple
    styles: dict
    properties: dict


def require_openpyxl():
    if not HAS_OPENPYXL:
        raise ReportError("openpyxl>=3.1,<4 must be installed in the Python running this module")
    version = tuple(int(part) for part in openpyxl.__version__.split(".")[:2])
    if not (3, 1) <= version < (4, 0):
        raise ReportError("openpyxl>=3.1,<4 is required; found %s" % openpyxl.__version__)


def _mapping(value, where, allowed=None):
    if not isinstance(value, dict):
        raise ReportError("%s must be a mapping" % where)
    if allowed is not None:
        unknown = set(value) - set(allowed)
        if unknown:
            raise ReportError("%s has unknown options: %s" % (where, ", ".join(map(str, unknown))))
    return value


def _sequence(value, where):
    if not isinstance(value, (list, tuple)):
        raise ReportError("%s must be a list" % where)
    return value


def _text(value, where, empty=False):
    if not isinstance(value, str) or (not empty and not value.strip()):
        raise ReportError("%s must be %s string" % (where, "a" if empty else "a non-empty"))
    if len(value) > 32767 or INVALID_XML.search(value):
        raise ReportError("%s contains invalid XML characters or exceeds 32767 characters" % where)
    return value


def _boolean(value, where):
    if not isinstance(value, bool):
        raise ReportError("%s must be a boolean" % where)
    return value


def _number(value, where, minimum=None, maximum=None):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ReportError("%s must be a finite number" % where)
    if isinstance(value, float) and not math.isfinite(value):
        raise ReportError("%s must be a finite number" % where)
    if (minimum is not None and value < minimum) or (maximum is not None and value > maximum):
        raise ReportError("%s is outside the supported range" % where)
    return value


def _color(value, where):
    value = _text(value, where).removeprefix("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}(?:[0-9A-F]{2})?", value):
        raise ReportError("%s must be a six-digit RGB or eight-digit ARGB color" % where)
    return "FF" + value if len(value) == 6 else value


def _style_name(value, styles, where):
    if value is None:
        return None
    name = _text(value, where)
    if name not in styles:
        raise ReportError("%s references undefined style %r" % (where, name))
    return name


def _range(value, where, max_column=MAX_COLUMNS, max_row=MAX_ROWS, cell=False):
    value = _text(value, where).replace("$", "").upper()
    pattern = CELL_REFERENCE if cell else RANGE_REFERENCE
    if not pattern.fullmatch(value):
        raise ReportError("%s must be an A1 %s" % (where, "cell reference" if cell else "range"))
    left, top, right, bottom = range_boundaries(value)
    if left > right or top > bottom or right > max_column or bottom > max_row:
        raise ReportError("%s is outside the worksheet bounds or has reversed endpoints" % where)
    return value


def _normalize_style(value, where):
    value = _mapping(
        value, where, {"font", "fill", "background", "border", "alignment", "number_format"}
    )
    result = {}
    if "font" in value:
        font = dict(
            _mapping(
                value["font"],
                where + ".font",
                {
                    "name",
                    "size",
                    "bold",
                    "italic",
                    "color",
                    "underline",
                    "strike",
                },
            )
        )
        for key in ("bold", "italic", "strike"):
            if key in font:
                _boolean(font[key], where + ".font." + key)
        if "size" in font:
            _number(font["size"], where + ".font.size", 1, 409)
        if "name" in font:
            _text(font["name"], where + ".font.name")
        if "color" in font:
            font["color"] = _color(font["color"], where + ".font.color")
        if "underline" in font and font["underline"] not in (
            "none",
            "single",
            "double",
            "singleAccounting",
            "doubleAccounting",
        ):
            raise ReportError(where + ".font.underline is invalid")
        result["font"] = font
    if "fill" in value and "background" in value:
        raise ReportError(where + " must use either fill or background")
    if "fill" in value:
        fill = _mapping(value["fill"], where + ".fill", {"color"})
        result["fill"] = {"color": _color(fill.get("color"), where + ".fill.color")}
    elif "background" in value:
        result["fill"] = {"color": _color(value["background"], where + ".background")}
    if "alignment" in value:
        alignment = value["alignment"]
        if isinstance(alignment, str):
            alignment = {"horizontal": alignment}
        alignment = dict(
            _mapping(
                alignment,
                where + ".alignment",
                {
                    "horizontal",
                    "vertical",
                    "wrap_text",
                    "shrink_to_fit",
                    "text_rotation",
                    "indent",
                },
            )
        )
        choices = {
            "horizontal": {
                "general",
                "left",
                "center",
                "right",
                "fill",
                "justify",
                "centerContinuous",
                "distributed",
            },
            "vertical": {"top", "center", "bottom", "justify", "distributed"},
        }
        for key, options in choices.items():
            if key in alignment and (
                not isinstance(alignment[key], str) or alignment[key] not in options
            ):
                raise ReportError(where + ".alignment." + key + " is invalid")
        for key in ("wrap_text", "shrink_to_fit"):
            if key in alignment:
                _boolean(alignment[key], where + ".alignment." + key)
        for key, maximum in (("text_rotation", 180), ("indent", 250)):
            if key in alignment:
                _number(alignment[key], where + ".alignment." + key, 0, maximum)
                if not isinstance(alignment[key], int):
                    raise ReportError(where + ".alignment." + key + " must be an integer")
        result["alignment"] = alignment
    if "border" in value:
        border = _mapping(value["border"], where + ".border", {"left", "right", "top", "bottom"})
        result["border"] = {}
        for edge, definition in border.items():
            edge_where = where + ".border." + edge
            side = dict(_mapping(definition, edge_where, {"style", "color"}))
            if side.get("style") not in (
                "dashDot",
                "dashDotDot",
                "dashed",
                "dotted",
                "double",
                "hair",
                "medium",
                "mediumDashDot",
                "mediumDashDotDot",
                "mediumDashed",
                "slantDashDot",
                "thick",
                "thin",
            ):
                raise ReportError(edge_where + ".style is invalid or missing")
            if "color" in side:
                side["color"] = _color(side["color"], edge_where + ".color")
            result["border"][edge] = side
    if "number_format" in value:
        result["number_format"] = _text(value["number_format"], where + ".number_format")
    if not result:
        raise ReportError(where + " must define at least one style component")
    return result


def _cell(value, styles, where, default_style=None, default_format=None):
    style, number_format, formula, cell_type = default_style, default_format, False, None
    if isinstance(value, dict):
        definition = _mapping(value, where, {"value", "formula", "type", "style", "number_format"})
        if ("value" in definition) == ("formula" in definition):
            raise ReportError(where + " must define exactly one of value or formula")
        style = _style_name(definition.get("style", style), styles, where + ".style")
        if "number_format" in definition:
            number_format = _text(definition["number_format"], where + ".number_format")
        formula = "formula" in definition
        cell_type = definition.get("type")
        if formula:
            if cell_type is not None:
                raise ReportError(where + " cannot combine formula with type")
            value = _text(definition["formula"], where + ".formula")
            if not value.startswith("=") or len(value) == 1:
                raise ReportError(where + ".formula must start with '=' and contain an expression")
        else:
            value = definition["value"]
    if cell_type is not None and not isinstance(cell_type, str):
        raise ReportError(where + ".type must be a string")
    if cell_type in ("date", "datetime"):
        try:
            if isinstance(value, str):
                value = (
                    date.fromisoformat(value)
                    if cell_type == "date"
                    else datetime.fromisoformat(value)
                )
            if cell_type == "date" and (not isinstance(value, date) or isinstance(value, datetime)):
                raise ValueError("expected a date")
            if cell_type == "datetime" and not isinstance(value, datetime):
                raise ValueError("expected a datetime")
        except (ValueError, TypeError) as exc:
            raise ReportError(where + " requires a valid ISO " + cell_type) from exc
    elif cell_type is not None:
        valid = {
            "string": isinstance(value, str),
            "number": isinstance(value, (int, float)) and not isinstance(value, bool),
            "boolean": isinstance(value, bool),
        }
        if not valid.get(cell_type, False):
            raise ReportError(
                where + " has an unsupported type or a value incompatible with its type"
            )
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            raise ReportError(where + " must use a datetime without a timezone")
    elif isinstance(value, str):
        _text(value, where, empty=True)
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        _number(value, where)
        # Excel stores numbers as doubles; avoid overflowing when serializing large integers.
        try:
            finite = math.isfinite(float(value))
        except OverflowError:
            finite = False
        if not finite:
            raise ReportError(where + " exceeds Excel's numeric range")
    elif value is not None and not isinstance(value, (bool, date)):
        raise ReportError(where + " has an unsupported cell value")
    return CellSpec(value, style, number_format, formula)


def _rows(values, columns, styles, where):
    values = _sequence(values, where)
    if len(values) >= MAX_ROWS:
        raise ReportError(where + " exceeds Excel's row limit including the header")
    names = {column.name for column in columns}
    rows = []
    for index, row in enumerate(values):
        row_where = "%s[%d]" % (where, index)
        row_style, cell_styles = None, {}
        if isinstance(row, dict):
            if set(row) <= names:
                row = [row.get(column.name) for column in columns]
            else:
                definition = _mapping(row, row_where, {"values", "style", "styles"})
                row_style = _style_name(definition.get("style"), styles, row_where + ".style")
                cell_styles = _mapping(definition.get("styles", {}), row_where + ".styles", names)
                cell_styles = {
                    key: _style_name(style, styles, row_where + ".styles." + key)
                    for key, style in cell_styles.items()
                }
                row = definition.get("values")
        row = _sequence(row, row_where + ".values")
        if len(row) != len(columns):
            raise ReportError("%s must contain exactly %d cells" % (row_where, len(columns)))
        rows.append(
            tuple(
                _cell(
                    value,
                    styles,
                    "%s[%d]" % (row_where, column_index),
                    cell_styles.get(column.name, row_style or column.style),
                    column.number_format,
                )
                for column_index, (column, value) in enumerate(zip(columns, row))
            )
        )
    return tuple(rows)


def _tables(values, columns, row_count, all_names, where):
    tables, occupied = [], []
    for index, value in enumerate(_sequence(values, where)):
        table_where = "%s[%d]" % (where, index)
        value = _mapping(
            value,
            table_where,
            {
                "name",
                "ref",
                "style",
                "show_row_stripes",
                "show_column_stripes",
                "show_first_column",
                "show_last_column",
            },
        )
        name = _text(value.get("name"), table_where + ".name")
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]{0,254}", name) or name.upper() in {"R", "C"}:
            raise ReportError(table_where + ".name must be a valid Excel table name")
        if CELL_REFERENCE.fullmatch(name.upper()):
            left, top, _, _ = range_boundaries(name.upper())
            if left <= MAX_COLUMNS and top <= MAX_ROWS:
                raise ReportError(table_where + ".name cannot be a cell reference")
        if re.fullmatch(r"R[1-9][0-9]*C[1-9][0-9]*", name.upper()):
            raise ReportError(table_where + ".name cannot be an R1C1 reference")
        if name.casefold() in all_names:
            raise ReportError(table_where + ".name duplicates another table name")
        all_names.add(name.casefold())
        ref = _range(
            value.get("ref", "A1:%s%d" % (get_column_letter(len(columns)), row_count)),
            table_where + ".ref",
            len(columns),
            row_count,
        )
        bounds = range_boundaries(ref)
        if bounds[1] != 1:
            raise ReportError(table_where + ".ref must include the column headers in row 1")
        for previous in occupied:
            if bounds[0] <= previous[2] and previous[0] <= bounds[2]:
                raise ReportError(table_where + ".ref overlaps another table")
        occupied.append(bounds)
        style = value.get("style", "TableStyleMedium2")
        match = re.fullmatch(
            r"TableStyle(Light|Medium|Dark)([1-9][0-9]*)", _text(style, table_where + ".style")
        )
        if not match or int(match[2]) > {"Light": 21, "Medium": 28, "Dark": 11}[match[1]]:
            raise ReportError(table_where + ".style must be a built-in Excel table style")
        table = {"name": name, "ref": ref, "style": style}
        for key in (
            "show_row_stripes",
            "show_column_stripes",
            "show_first_column",
            "show_last_column",
        ):
            table[key] = _boolean(
                value.get(key, key == "show_row_stripes"), table_where + "." + key
            )
        tables.append(table)
    return tuple(tables)


def _auto_filter(value, columns, row_count, tables, where):
    if isinstance(value, bool):
        # Tables already provide filters. A worksheet filter over the same cells
        # makes Excel repair the workbook by removing the table.
        return (
            "A1:%s%d" % (get_column_letter(len(columns)), row_count)
            if value and not tables
            else None
        )
    ref = _range(value, where, len(columns), row_count)
    bounds = range_boundaries(ref)
    for table in tables:
        table_bounds = range_boundaries(table["ref"])
        if bounds == table_bounds:
            return None
        if (
            bounds[0] <= table_bounds[2]
            and table_bounds[0] <= bounds[2]
            and bounds[1] <= table_bounds[3]
            and table_bounds[1] <= bounds[3]
        ):
            raise ReportError(where + " cannot overlap a table; tables provide their own filters")
    return ref


def _formats(values, columns, row_count, styles, where):
    rules = []
    column_indexes = {column.name: index for index, column in enumerate(columns, 1)}
    for index, value in enumerate(_sequence(values, where)):
        rule_where = "%s[%d]" % (where, index)
        value = _mapping(
            value,
            rule_where,
            {
                "column",
                "range",
                "equals",
                "formula",
                "style",
                "background",
                "font_color",
                "stop_if_true",
            },
        )
        if ("column" in value) == ("range" in value):
            raise ReportError(rule_where + " must define exactly one of column or range")
        if ("equals" in value) == ("formula" in value):
            raise ReportError(rule_where + " must define exactly one of equals or formula")
        if "column" in value:
            column = _text(value["column"], rule_where + ".column")
            if column not in column_indexes:
                raise ReportError(rule_where + ".column references an undefined column")
            letter = get_column_letter(column_indexes[column])
            ref = "%s2:%s%d" % (letter, letter, row_count) if row_count > 1 else None
        else:
            ref = _range(value["range"], rule_where + ".range", len(columns), row_count)
        if "equals" in value:
            if not isinstance(value["equals"], (str, int, float, bool)):
                raise ReportError(rule_where + ".equals must be a string, number, or boolean")
            literal = _cell(value["equals"], styles, rule_where + ".equals").value
            if isinstance(literal, str):
                formula = '"%s"' % literal.replace('"', '""')
            elif isinstance(literal, bool):
                formula = "TRUE" if literal else "FALSE"
            elif isinstance(literal, (int, float)):
                formula = str(literal)
            else:
                raise ReportError(rule_where + ".equals must be a string, number, or boolean")
            kind = "cellIs"
        else:
            formula = _text(value["formula"], rule_where + ".formula").removeprefix("=")
            if not formula.strip():
                raise ReportError(rule_where + ".formula must contain an expression")
            kind = "expression"
        rule = {"ref": ref, "kind": kind, "formula": formula}
        rule["style"] = _style_name(value.get("style"), styles, rule_where + ".style")
        for key in ("background", "font_color"):
            rule[key] = _color(value[key], rule_where + "." + key) if key in value else None
        rule["stop_if_true"] = _boolean(
            value.get("stop_if_true", False), rule_where + ".stop_if_true"
        )
        if not any(rule[key] for key in ("style", "background", "font_color")):
            raise ReportError(rule_where + " must define style, background, or font_color")
        if ref is not None:
            rules.append(rule)
    return tuple(rules)


def normalize_report(sheets, styles=None, properties=None):
    """Return a validated model; never mutate the caller's dictionaries."""
    require_openpyxl()
    styles = _mapping(styles if styles is not None else {}, "styles")
    normalized_styles, style_names = {}, set()
    for name, definition in styles.items():
        _text(name, "style name")
        if name.casefold() == "normal":
            raise ReportError("styles cannot redefine the built-in Normal style")
        if name.casefold() in style_names:
            raise ReportError("styles contains duplicate names without regard to case")
        style_names.add(name.casefold())
        normalized_styles[name] = _normalize_style(definition, "styles." + name)
    properties = _mapping(
        properties if properties is not None else {},
        "properties",
        {
            "title",
            "subject",
            "creator",
            "description",
            "keywords",
            "category",
            "language",
            "last_modified_by",
        },
    )
    properties = {
        name: _text(value, "properties." + name, empty=True) for name, value in properties.items()
    }
    sheets = _sequence(sheets, "sheets")
    if not sheets:
        raise ReportError("sheets must contain at least one worksheet")
    normalized_sheets, sheet_names, table_names = [], set(), set()
    for index, definition in enumerate(sheets):
        where = "sheets[%d]" % index
        definition = _mapping(
            definition,
            where,
            {
                "name",
                "columns",
                "rows",
                "header_style",
                "freeze_panes",
                "auto_filter",
                "tables",
                "conditional_formatting",
            },
        )
        name = _text(definition.get("name"), where + ".name")
        if (
            len(name) > 31
            or re.search(r"[\\/*?:\[\]]", name)
            or name.startswith("'")
            or name.endswith("'")
        ):
            raise ReportError(where + ".name is not a valid Excel worksheet name")
        if name.casefold() in sheet_names:
            raise ReportError(where + ".name duplicates another worksheet name")
        sheet_names.add(name.casefold())
        header_style = _style_name(
            definition.get("header_style"), normalized_styles, where + ".header_style"
        )
        columns, column_names = [], set()
        for column_index, column in enumerate(
            _sequence(definition.get("columns"), where + ".columns")
        ):
            column_where = "%s.columns[%d]" % (where, column_index)
            column = _mapping(
                column, column_where, {"name", "width", "style", "header_style", "number_format"}
            )
            column_name = _text(column.get("name"), column_where + ".name")
            if column_name.casefold() in column_names:
                raise ReportError(column_where + ".name duplicates another column name")
            column_names.add(column_name.casefold())
            width = (
                _number(column["width"], column_where + ".width", 0.1, 255)
                if "width" in column
                else None
            )
            number_format = (
                _text(column["number_format"], column_where + ".number_format")
                if "number_format" in column
                else None
            )
            columns.append(
                ColumnSpec(
                    column_name,
                    width,
                    _style_name(column.get("style"), normalized_styles, column_where + ".style"),
                    _style_name(
                        column.get("header_style", header_style),
                        normalized_styles,
                        column_where + ".header_style",
                    ),
                    number_format,
                )
            )
        if not columns or len(columns) > MAX_COLUMNS:
            raise ReportError(where + ".columns must contain between 1 and 16384 columns")
        rows = _rows(definition.get("rows", []), columns, normalized_styles, where + ".rows")
        freeze_panes = definition.get("freeze_panes")
        if freeze_panes is not None:
            freeze_panes = _range(freeze_panes, where + ".freeze_panes", cell=True)
            if freeze_panes == "A1":
                freeze_panes = None
        tables = _tables(
            definition.get("tables", []), columns, len(rows) + 1, table_names, where + ".tables"
        )
        auto_filter = _auto_filter(
            definition.get("auto_filter", False),
            columns,
            len(rows) + 1,
            tables,
            where + ".auto_filter",
        )
        normalized_sheets.append(
            SheetSpec(
                name,
                tuple(columns),
                rows,
                freeze_panes,
                auto_filter,
                tables,
                _formats(
                    definition.get("conditional_formatting", []),
                    columns,
                    len(rows) + 1,
                    normalized_styles,
                    where + ".conditional_formatting",
                ),
            )
        )
    return WorkbookSpec(tuple(normalized_sheets), normalized_styles, properties)


def _style_components(definition):
    components = {}
    if "font" in definition:
        components["font"] = Font(**definition["font"])
    if "fill" in definition:
        components["fill"] = PatternFill(fill_type="solid", fgColor=definition["fill"]["color"])
    if "alignment" in definition:
        components["alignment"] = Alignment(**definition["alignment"])
    if "border" in definition:
        components["border"] = Border(
            **{edge: Side(**side) for edge, side in definition["border"].items()}
        )
    return components


def render_workbook(spec):
    """Render the model to deterministic XLSX bytes, without touching the filesystem."""
    require_openpyxl()
    workbook = Workbook(iso_dates=True)
    workbook.remove(workbook.active)
    workbook.properties.created = EPOCH
    workbook.properties.modified = EPOCH
    workbook.properties.creator = "ansible_tools.reporting"
    for key, value in spec.properties.items():
        setattr(workbook.properties, "lastModifiedBy" if key == "last_modified_by" else key, value)
    for name in sorted(spec.styles):
        definition = spec.styles[name]
        style = NamedStyle(name=name, **_style_components(definition))
        if "number_format" in definition:
            style.number_format = definition["number_format"]
        workbook.add_named_style(style)
    for sheet in spec.sheets:
        worksheet = workbook.create_sheet(sheet.name)
        for index, column in enumerate(sheet.columns, 1):
            header = worksheet.cell(1, index, column.name)
            header.data_type = "s"
            if column.header_style:
                header.style = column.header_style
            if column.width is not None:
                worksheet.column_dimensions[get_column_letter(index)].width = column.width
        for row_index, row in enumerate(sheet.rows, 2):
            for column_index, definition in enumerate(row, 1):
                cell = worksheet.cell(row_index, column_index, definition.value)
                inferred_format = cell.number_format
                if isinstance(definition.value, str):
                    cell.data_type = "f" if definition.formula else "s"
                if definition.style:
                    cell.style = definition.style
                    if (
                        isinstance(definition.value, date)
                        and "number_format" not in spec.styles[definition.style]
                    ):
                        cell.number_format = inferred_format
                if definition.number_format:
                    cell.number_format = definition.number_format
        worksheet.freeze_panes = sheet.freeze_panes
        worksheet.auto_filter.ref = sheet.auto_filter
        for definition in sheet.tables:
            table = Table(displayName=definition["name"], ref=definition["ref"])
            table.tableStyleInfo = TableStyleInfo(
                name=definition["style"],
                showRowStripes=definition["show_row_stripes"],
                showColumnStripes=definition["show_column_stripes"],
                showFirstColumn=definition["show_first_column"],
                showLastColumn=definition["show_last_column"],
            )
            worksheet.add_table(table)
        for definition in sheet.conditional_formatting:
            style = spec.styles.get(definition["style"], {})
            components = _style_components(style)
            if "number_format" in style:
                components["numFmt"] = NumberFormat(numFmtId=164, formatCode=style["number_format"])
            if definition["background"]:
                components["fill"] = PatternFill(
                    fill_type="solid", fgColor=definition["background"]
                )
            if definition["font_color"]:
                font = dict(style.get("font", {}), color=definition["font_color"])
                components["font"] = Font(**font)
            rule = Rule(
                type=definition["kind"],
                formula=[definition["formula"]],
                dxf=DifferentialStyle(**components),
                stopIfTrue=definition["stop_if_true"],
                operator="equal" if definition["kind"] == "cellIs" else None,
            )
            worksheet.conditional_formatting.add(definition["ref"], rule)
    stream = BytesIO()
    try:
        workbook.save(stream)
    finally:
        workbook.close()
    result = BytesIO()
    with ZipFile(stream) as source, ZipFile(result, "w", compression=ZIP_DEFLATED) as output:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            output.writestr(entry, _normalized_part(name, source.read(name)))
    return result.getvalue()


def _normalized_part(name, data):
    if name != "docProps/core.xml":
        return data
    # Keep the dcterms prefix bound: xsi:type contains the QName dcterms:W3CDTF.
    for prefix, uri in (
        ("cp", "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"),
        ("dc", "http://purl.org/dc/elements/1.1/"),
        ("dcterms", "http://purl.org/dc/terms/"),
        ("xsi", "http://www.w3.org/2001/XMLSchema-instance"),
    ):
        ElementTree.register_namespace(prefix, uri)
    root = ElementTree.fromstring(data)
    namespace = "{http://purl.org/dc/terms/}"
    for field in ("created", "modified"):
        element = root.find(namespace + field)
        if element is not None:
            element.text = "2000-01-01T00:00:00Z"
    return ElementTree.tostring(root, encoding="utf-8")


def workbook_checksum(data):
    """Hash package contents, ignoring ZIP metadata and volatile save timestamps.

    Return None for an invalid package so overwrite can repair a corrupt output.
    All other package contents participate; hand-edited reports are regenerated.
    """
    digest = hashlib.sha256()
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = archive.namelist()
            if len(names) != len(set(names)) or not {
                "[Content_Types].xml",
                "xl/workbook.xml",
            } <= set(names):
                return None
            for name in sorted(names):
                payload = _normalized_part(name, archive.read(name))
                encoded_name = name.encode("utf-8")
                digest.update(len(encoded_name).to_bytes(8, "big"))
                digest.update(encoded_name)
                digest.update(len(payload).to_bytes(8, "big"))
                digest.update(payload)
    except (BadZipFile, ElementTree.ParseError, KeyError, ValueError, RuntimeError, zlib.error):
        return None
    return digest.hexdigest()
